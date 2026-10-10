"""Regression tests for the PowerPoint scripts, run as real child processes.

- audit_deck.py and bulk_read.py: an 11-slide deck is BUILT at test time with python-pptx (no binary .pptx is
  committed) and read through PowerPoint; the output is compared with fixtures/expected. These run only where
  python-pptx, pywin32 and PowerPoint are installed, and quit PowerPoint afterwards if the tests started it.
- contact_sheet.py --renders: a sheet from PNGs generated with pillow (any OS).

Run: python -m unittest discover -s powerpoint/tests -v
KSHAREDPY_UPDATE_EXPECTED=1 rewrites fixtures/expected (review the diff before committing).
"""
import importlib.util
import os
import sys
import unittest

TESTS = os.path.dirname(os.path.abspath(__file__))
POWERPOINT = os.path.dirname(TESTS)
EXPECTED = os.path.join(TESTS, "fixtures", "expected")
sys.path.insert(0, os.path.join(os.path.dirname(POWERPOINT), "kshared", "tests", "support"))
import kScriptHarness as H  # noqa: E402

COM_TRANSPORT_ERRORS = ("The RPC server is unavailable", "The remote procedure call failed",
                        "-2147023174", "-2147023170", "-2147417848")  # the PowerPoint process went away
HAS_PIL = importlib.util.find_spec("PIL") is not None
HAS_PPTX = importlib.util.find_spec("pptx") is not None

# (title, body) per slide: cover/agenda/rules slides skipped by number, a tight and a long slide, a font below
# the floor, an opponent (HOPP) slide and a vote slide with non-ASCII text.
DECK = [("Cover", "x"), ("Hosts", "a b"), ("Agenda", "c"), ("Rules", "d"), ("Five", "e"),
        ("Short slide", "three words here"), ("Tight one", "one two three four five six seven eight nine ten"),
        ("Long slide", "w " * 20), ("Small font", "tiny text"), ("HOPP intro", "skip me please"),
        ("Vote", "Rösta nu → CO₂ 44 °")]


def HasPowerPoint():
    if sys.platform != "win32" or importlib.util.find_spec("win32com") is None:
        return False
    import winreg
    try:
        winreg.CloseKey(winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, "PowerPoint.Application"))
        return True
    except OSError:
        return False


def BuildDeck(Path):
    """The 11-slide test deck, written with python-pptx."""
    from pptx import Presentation
    from pptx.util import Pt
    Deck = Presentation()
    for Index, (Title, Body) in enumerate(DECK):
        Slide = Deck.slides.add_slide(Deck.slide_layouts[1])
        Slide.shapes.title.text = Title
        Slide.placeholders[1].text = Body
        for Paragraph in Slide.placeholders[1].text_frame.paragraphs:
            for Run in Paragraph.runs:
                Run.font.size = Pt(12 if Title == "Small font" else 24)
        Slide.notes_slide.notes_text_frame.text = f"notes {Index + 1} ₂"
    Deck.save(Path)


class kPowerPointHold:
    """Keeps one PowerPoint instance alive for a test class. Without it every child script starts PowerPoint and
    quits it, and the next child can attach to the instance that is still shutting down ("The RPC server is
    unavailable"). With it, kPptSession in the child finds a running instance and leaves it running."""

    def __init__(self):
        import pythoncom
        import time
        import win32com.client
        pythoncom.CoInitialize()
        try:
            win32com.client.GetActiveObject("PowerPoint.Application")
            self.WasRunning = True
        except pythoncom.com_error:  # not running: this hold starts it and must quit it
            self.WasRunning = False
        self.App = None
        for _ in range(20):  # an instance that a previous test is quitting refuses calls for a moment
            try:
                App = win32com.client.Dispatch("PowerPoint.Application")
                App.Version
                self.App = App
                break
            except pythoncom.com_error:  # the dying instance: wait for it to go
                time.sleep(0.5)

    def Release(self):
        import pythoncom
        try:
            if self.App is not None and not self.WasRunning and self.App.Presentations.Count == 0:
                self.App.Quit()
        except pythoncom.com_error:  # already gone: nothing to close
            pass
        self.App = None
        pythoncom.CoUninitialize()


@unittest.skipUnless(HAS_PPTX and HasPowerPoint(), "needs python-pptx, pywin32 and PowerPoint")
class kDeckScriptTests(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.Work = H.NewWorkFolder("ppt-scripts")
        cls.Deck = os.path.join(cls.Work, "test.pptx")
        BuildDeck(cls.Deck)
        cls.Hold = kPowerPointHold()

    @classmethod
    def tearDownClass(cls):
        cls.Hold.Release()
        H.RemoveWorkFolder(cls.Work)

    def Run(self, Script, *Args):
        """Run a script; a run killed by PowerPoint going away under it (another process on the machine quitting
        the shared instance) is retried, up to 3 attempts - any other failure is not."""
        for _ in range(3):
            Code, Out, Err = H.Run(os.path.join(POWERPOINT, Script), list(Args), self.Work, H.BaseEnv())
            if Code == 0 or not any(Text in Err for Text in COM_TRANSPORT_ERRORS):
                break
            type(self).Hold.Release()
            type(self).Hold = kPowerPointHold()
        self.assertEqual(Code, 0, Err)
        return Out, Err

    def test_AuditTable(self):
        Out, _ = self.Run("audit_deck.py", "--file", self.Deck)
        H.AssertMatches(self, Out, os.path.join(EXPECTED, "audit_default.txt"))

    def test_AuditAllWithExtraSkip(self):
        Out, _ = self.Run("audit_deck.py", "--file", self.Deck, "--all", "--skip", "Tight")
        H.AssertMatches(self, Out, os.path.join(EXPECTED, "audit_all_skip.txt"))

    def test_AuditJson(self):
        Out, _ = self.Run("audit_deck.py", "--file", self.Deck, "--json")
        H.AssertMatches(self, Out, os.path.join(EXPECTED, "audit_json.json"))

    def test_BulkReadAll(self):
        Out, _ = self.Run("bulk_read.py", "--file", self.Deck)
        H.AssertMatches(self, Out, os.path.join(EXPECTED, "bulk_default.json"))

    def test_BulkReadRangeWithoutNotes(self):
        Out, _ = self.Run("bulk_read.py", "--file", self.Deck, "--start", "6", "--end", "8", "--no-notes", "--pretty")
        H.AssertMatches(self, Out, os.path.join(EXPECTED, "bulk_range.json"))

    def test_BulkReadToFile(self):
        Target = os.path.join(self.Work, "bulk.json")
        Out, Err = self.Run("bulk_read.py", "--file", self.Deck, "--out", Target)
        self.assertEqual(Out, "")
        self.assertIn(f"Wrote 11 slides to {Target}", Err)
        H.AssertMatches(self, H.ReadText(Target), os.path.join(EXPECTED, "bulk_out.json"))

    def test_MissingDeckIsBadInput(self):
        Code, _, Err = H.Run(os.path.join(POWERPOINT, "bulk_read.py"), ["--file", "absent.pptx"], self.Work,
                             H.BaseEnv())
        self.assertEqual(Code, 2, Err)
        self.assertIn("File not found", Err)


@unittest.skipUnless(HAS_PIL, "needs pillow")
class kContactSheetTests(unittest.TestCase):
    """contact_sheet.py --renders on PNGs made here. Layout: PAD 8, caption 18."""

    COLOURS = [(255, 0, 0), (0, 160, 0), (0, 0, 255), (200, 200, 0), (0, 200, 200), (128, 0, 128), (60, 60, 60)]

    def setUp(self):
        from PIL import Image
        self.Work = H.NewWorkFolder("contact-sheet")
        self.Renders = os.path.join(self.Work, "renders")
        os.makedirs(self.Renders)
        # Slide numbers out of name order (s10 after s9), a stale contact sheet and a non-image to ignore.
        for Number, Colour in zip((1, 2, 3, 9, 10, 11, 12), self.COLOURS):
            Image.new("RGB", (320, 180), Colour).save(os.path.join(self.Renders, f"s{Number}.png"))
        Image.new("RGB", (50, 50), (0, 0, 0)).save(os.path.join(self.Renders, "old_contact.png"))
        with open(os.path.join(self.Renders, "notes.txt"), "w", encoding="utf-8") as File:
            File.write("not an image")

    def tearDown(self):
        H.RemoveWorkFolder(self.Work)

    def Run(self, *Args):
        return H.Run(os.path.join(POWERPOINT, "contact_sheet.py"), list(Args), self.Work, H.BaseEnv())

    def test_SheetLayoutAndOrder(self):
        from PIL import Image
        Out = os.path.join(self.Work, "sheet.png")
        Code, Stdout, Err = self.Run("--renders", self.Renders, "--cols", "3", "--slide-width", "100", "--out", Out)
        self.assertEqual(Code, 0, Err)
        self.assertEqual(Stdout.strip(), Out)
        with Image.open(Out) as Sheet:
            Sheet = Sheet.convert("RGB")
            Width, Height = 100, 56  # 100 * 180 / 320, rounded
            self.assertEqual(Sheet.size, (3 * (Width + 8) + 8, 3 * (Height + 18 + 8) + 8))
            for Index, Colour in enumerate(self.COLOURS):
                X = 8 + (Index % 3) * (Width + 8) + Width // 2
                Y = 8 + (Index // 3) * (Height + 18 + 8) + Height // 2
                self.assertEqual(Sheet.getpixel((X, Y)), Colour, f"tile {Index}")
            self.assertEqual(Sheet.getpixel((8 + 2 * (Width + 8) + Width // 2, 8 + 2 * (Height + 26) + Height // 2)),
                             (255, 255, 255), "the empty cells stay white")

    def test_DefaultOutputBesideRenders(self):
        Code, Stdout, Err = self.Run("--renders", self.Renders)
        self.assertEqual(Code, 0, Err)
        Expected = os.path.join(self.Renders, "contact.png")
        self.assertEqual(Stdout.strip(), Expected)
        self.assertTrue(os.path.isfile(Expected))

    def test_MissingFolderIsBadInput(self):
        Code, _, Err = self.Run("--renders", os.path.join(self.Work, "nothere"))
        self.assertEqual(Code, 2, Err)
        self.assertIn("renders folder not found", Err)

    def test_EmptyFolderIsReportedFailure(self):
        Empty = os.path.join(self.Work, "empty")
        os.makedirs(Empty)
        Code, _, Err = self.Run("--renders", Empty)
        self.assertEqual(Code, 1, Err)
        self.assertIn("no slides rendered", Err)


if __name__ == "__main__":
    unittest.main()
