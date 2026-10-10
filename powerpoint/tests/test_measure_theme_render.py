"""Tests for _measure.py, _theme.py and render_lo.py.

Run: python -m unittest discover -s X:\\Source\\Admin\\kSharedPy\\powerpoint\\tests -v

- _measure: the wrap arithmetic with no fonts at all (FONT_DIRS emptied: the documented 0.5-em fallback makes the
  widths exact), and with real metrics where Carlito or Liberation Sans is installed (CI installs them).
- _theme: a deck built at test time with python-pptx (skipped without it); colour adjustments on hand-made XML.
- render_lo: the soffice lookup and page ordering everywhere; a real render only where LibreOffice and pdftoppm are.
Work folders live under kSharedPy\\_scratch\\render-tests, never in the user profile.
"""
import importlib.util
import os
import shutil
import sys
import unittest
import uuid

POWERPOINT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SCRATCH = os.path.join(os.path.dirname(POWERPOINT), "_scratch", "render-tests")
sys.path.insert(0, POWERPOINT)

from kShared import ToolReportableException, kErrorReport, kS  # noqa: E402
from render_lo import kLoRenderer  # noqa: E402

HAS_PIL = importlib.util.find_spec("PIL") is not None
HAS_PPTX = importlib.util.find_spec("pptx") is not None and importlib.util.find_spec("lxml") is not None
A = "http://schemas.openxmlformats.org/drawingml/2006/main"


class kTestBase(unittest.TestCase):

    def setUp(self):
        kS.Reset()
        kErrorReport.Reset()
        os.environ["KPS_ERROR_REPORT"] = "0"

    def tearDown(self):
        self.assertFalse(kS.ErrorMode, "a master function reported an error")


@unittest.skipUnless(HAS_PIL, "pillow not installed")
class kMeasureFallbackTests(kTestBase):
    """No fonts: every glyph is 0.5 em x the family's width factor, so the arithmetic is exact."""

    def setUp(self):
        super().setUp()
        import _measure
        self.Module = _measure
        self.SavedDirs = _measure.FONT_DIRS
        _measure.FONT_DIRS = []
        self.ClearCaches()

    def tearDown(self):
        self.Module.FONT_DIRS = self.SavedDirs
        self.ClearCaches()
        super().tearDown()

    def ClearCaches(self):
        self.Module.kMeasure.Index.cache_clear()
        self.Module.kMeasure.Font.cache_clear()
        self.Module.kMeasure.Substitutes.clear()

    def test_TextWidthFallback(self):
        M = self.Module.kMeasure
        self.assertEqual(M.Index(), {})
        self.assertAlmostEqual(M.TextWidth("abcd", "Nofont", 10), 20.0)       # 4 x 0.5 x 10
        self.assertAlmostEqual(M.TextWidth("abcd", "Verdana", 10), 20.0 * 1.17)
        self.assertFalse(M.Installed("Calibri"))

    def test_WrapAndLineWords(self):
        M = self.Module.kMeasure
        # "aaaa" = 20 pt, a space = 5 pt: two words fit in 45 pt, three do not
        self.assertEqual(M.Wrap("aaaa aaaa aaaa", "Nofont", 10, 45), (2, 20.0))
        self.assertEqual(M.LineWords("aaaa aaaa aaaa", "Nofont", 10, 45), [["aaaa", "aaaa"], ["aaaa"]])
        self.assertEqual(M.Wrap("one\ntwo", "Nofont", 10, 1000)[0], 2)
        self.assertEqual(M.Wrap("a" * 30, "Nofont", 10, 50)[0], 3)            # 150 pt breaks inside the word
        self.assertEqual(M.LooseLines("aaaa aaaa aaaa", "Nofont", 10, 45), 2)

    def test_TextHeightAndSafeWidth(self):
        M = self.Module.kMeasure
        Height, Widest, Lines = M.TextHeight([("aaaa aaaa aaaa", "Nofont", 10, False, 0),
                                              ("aaaa", "Nofont", 10, False, 6)], 45)
        self.assertEqual(Lines, 3)
        self.assertAlmostEqual(Height, 3 * 10 * self.Module.LINE_HEIGHT + 6)
        self.assertEqual(Widest, 20.0)
        self.assertEqual(M.LooseHeight([("aaaa", "Nofont", 10, False, 0)], 45), M.TextHeight(
            [("aaaa", "Nofont", 10, False, 0)], 45))
        self.assertAlmostEqual(M.SafeWidth("ab", "Nofont", 10), 10 * 1.04)
        self.assertAlmostEqual(M.SafeWidth("a\u2192", "Nofont", 10), 10 * 1.04 + 2.0)  # an arrow adds 0.2 em

    def test_FontDirsVariableIsASetting(self):
        self.assertTrue(self.Module.FONT_DIRS_VARIABLE)


@unittest.skipUnless(HAS_PIL, "pillow not installed")
class kMeasureRealFontTests(kTestBase):
    """With real metrics: Calibri measures through its metric twin Carlito, Arial through Liberation Sans."""

    def test_MetricTwins(self):
        import _measure
        M = _measure.kMeasure
        Pairs = [(Family, Twin) for Family, Twin in (("Calibri", "carlito"), ("Arial", "liberation sans"))
                 if Twin in M.Index() or Family.lower() in M.Index()]
        if not Pairs:
            self.skipTest("no Calibri/Carlito or Arial/Liberation Sans installed (CI: fonts-crosextra-carlito)")
        for Family, Twin in Pairs:
            with self.subTest(family=Family):
                self.assertTrue(M.Installed(Family))
                Wide, Narrow = M.TextWidth("WWWW", Family, 12), M.TextWidth("iiii", Family, 12)
                self.assertGreater(Wide, Narrow * 2)
                if Family.lower() not in M.Index():  # measured through the twin: the same widths exactly
                    self.assertAlmostEqual(M.TextWidth("Hello", Family, 12), M.TextWidth("Hello", Twin, 12))
                self.assertGreater(M.TextWidth("Hello", Family, 12, Bold=True), 0)


@unittest.skipUnless(HAS_PPTX, "python-pptx not installed")
class kThemeTests(kTestBase):

    def setUp(self):
        super().setUp()
        from pptx import Presentation
        self.Pres = Presentation()
        from _theme import kTheme
        self.kTheme = kTheme
        self.Theme = kTheme(self.Pres.slide_master)

    def El(self, Xml):
        from lxml import etree
        return etree.fromstring(f'<a:root xmlns:a="{A}">{Xml}</a:root>')[0]

    def test_DefaultTheme(self):
        Theme = self.Theme
        self.assertEqual(Theme.Major, "Calibri")
        self.assertEqual(Theme.Minor, "Calibri")
        self.assertEqual(Theme.Font("+mj-lt"), "Calibri")
        self.assertEqual(Theme.Font("Arial"), "Arial")
        self.assertEqual(len(Theme.Colours), 12)
        self.assertEqual(Theme.Map["bg1"], "lt1")
        self.assertTrue(Theme.IsThemeHex(Theme.Colours["accent1"].lower()))

    def test_ResolveAndBackground(self):
        Theme = self.Theme
        self.assertEqual(Theme.Resolve(self.El('<a:srgbClr val="ff0000"/>')), "FF0000")
        self.assertEqual(Theme.Resolve(self.El('<a:schemeClr val="tx1"/>')), Theme.Colours["dk1"])
        self.assertEqual(Theme.Resolve(self.El('<a:prstClr val="white"/>')), "FFFFFF")
        self.assertIsNone(Theme.Resolve(self.El('<a:hslClr hue="0" sat="0" lum="0"/>')))
        self.assertEqual(Theme.ColourIn(self.El('<a:solidFill><a:srgbClr val="00FF00"/></a:solidFill>')), "00FF00")
        Slide = self.Pres.slides.add_slide(self.Pres.slide_layouts[6])
        self.assertEqual(Theme.Background(Slide), Theme.Colours["lt1"])
        Slide.background.fill.solid()
        from pptx.dml.color import RGBColor
        Slide.background.fill.fore_color.rgb = RGBColor(0x12, 0x34, 0x56)
        self.assertEqual(Theme.Background(Slide), "123456")

    def test_Adjustments(self):
        T = self.kTheme
        self.assertEqual(T.Adjust("808080", self.El('<a:srgbClr val="808080"><a:shade val="50000"/></a:srgbClr>')),
                         "404040")
        self.assertEqual(T.Adjust("000000", self.El('<a:srgbClr val="000000"><a:tint val="50000"/></a:srgbClr>')),
                         "808080")
        self.assertEqual(T.Adjust("FF0000", self.El(
            '<a:srgbClr val="FF0000"><a:lumMod val="50000"/></a:srgbClr>')), "800000")
        self.assertEqual(T.RgbToHex(*T.HexToRgb("1A2B3C")), "1A2B3C")
        self.assertAlmostEqual(T.Saturation("FF0000"), 1.0)
        self.assertAlmostEqual(T.Saturation("808080"), 0.0)
        self.assertAlmostEqual(T.Hue("00FF00"), 1 / 3)


class kLoRendererTests(kTestBase):

    def setUp(self):
        super().setUp()
        self.Saved = os.environ.pop("KPS_SOFFICE", None)
        self.Root = os.path.join(SCRATCH, uuid.uuid4().hex[:8])
        os.makedirs(self.Root)

    def tearDown(self):
        if self.Saved is not None:
            os.environ["KPS_SOFFICE"] = self.Saved
        else:
            os.environ.pop("KPS_SOFFICE", None)
        shutil.rmtree(self.Root, ignore_errors=True)
        super().tearDown()

    def test_PageNumber(self):
        Pages = ["page-10.png", "page-2.png", "page-1.png"]
        self.assertEqual(sorted(Pages, key=kLoRenderer.PageNumber), ["page-1.png", "page-2.png", "page-10.png"])

    def test_KpsSofficeFileOrFolder(self):
        Exe = os.path.join(self.Root, "soffice.exe")
        with open(Exe, "w", encoding="utf-8") as File:
            File.write("")
        os.environ["KPS_SOFFICE"] = Exe
        self.assertEqual(kLoRenderer.FindSoffice(), Exe)
        os.environ["KPS_SOFFICE"] = self.Root
        self.assertEqual(kLoRenderer.FindSoffice(), Exe)

    def test_KpsSofficeWrongIsReported(self):
        os.environ["KPS_SOFFICE"] = os.path.join(self.Root, "nothing-here")
        with self.assertRaises(ToolReportableException):
            kLoRenderer.FindSoffice()

    @unittest.skipUnless(HAS_PPTX, "python-pptx not installed")
    def test_VisibleSlidesSkipsHidden(self):
        from pptx import Presentation
        Pres = Presentation()
        for _ in range(3):
            Pres.slides.add_slide(Pres.slide_layouts[6])
        Pres.slides[1]._element.set("show", "0")
        Deck = os.path.join(self.Root, "deck.pptx")
        Pres.save(Deck)
        self.assertEqual(kLoRenderer.VisibleSlides(Deck), [1, 3])

    @unittest.skipUnless(HAS_PPTX, "python-pptx not installed")
    def test_RenderWhereLibreOfficeIs(self):
        if not kLoRenderer.Available():
            self.skipTest("LibreOffice or pdftoppm not installed")
        from pptx import Presentation
        Pres = Presentation()
        for Title in ("One", "Two"):
            Pres.slides.add_slide(Pres.slide_layouts[0]).shapes.title.text = Title
        Deck = os.path.join(self.Root, "deck.pptx")
        Pres.save(Deck)
        Files = kLoRenderer.Render(Deck, os.path.join(self.Root, "out"), 320)
        self.assertEqual([os.path.basename(F) for F in Files], ["s001.png", "s002.png"])


if __name__ == "__main__":
    unittest.main()
